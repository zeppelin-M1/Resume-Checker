"""
Resume text extraction.

Supports .pdf, .docx, and .txt uploads. Everything is done in-memory
(no files written to disk) using BytesIO.
"""

import io
from pypdf import PdfReader
from docx import Document
from fastapi import UploadFile, HTTPException

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}
MAX_FILE_SIZE_MB = 10


def _get_extension(filename: str) -> str:
    if "." not in filename:
        return ""
    return "." + filename.rsplit(".", 1)[-1].lower()


def _extract_pdf(file_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(file_bytes))
    text_parts = []
    for page in reader.pages:
        extracted = page.extract_text() or ""
        text_parts.append(extracted)
    return "\n".join(text_parts).strip()


def _extract_docx(file_bytes: bytes) -> str:
    doc = Document(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in doc.paragraphs]

    # Also pull text out of tables (many resumes use table layouts)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text)

    return "\n".join(paragraphs).strip()


def _extract_txt(file_bytes: bytes) -> str:
    return file_bytes.decode("utf-8", errors="ignore").strip()


async def extract_text_from_upload(file: UploadFile) -> str:
    """
    Reads an UploadFile, validates it, and returns extracted plain text.
    Raises HTTPException on invalid file type, empty file, or size limit.
    """
    ext = _get_extension(file.filename or "")
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail=f"Unsupported file type '{ext}'. Please upload PDF, DOCX, or TXT.",
        )

    file_bytes = await file.read()

    size_mb = len(file_bytes) / (1024 * 1024)
    if size_mb > MAX_FILE_SIZE_MB:
        raise HTTPException(status_code=400, detail=f"File too large ({size_mb:.1f}MB). Max {MAX_FILE_SIZE_MB}MB.")

    if len(file_bytes) == 0:
        raise HTTPException(status_code=400, detail="Uploaded file is empty.")

    try:
        if ext == ".pdf":
            text = _extract_pdf(file_bytes)
        elif ext == ".docx":
            text = _extract_docx(file_bytes)
        else:
            text = _extract_txt(file_bytes)
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"Could not parse file: {exc}")

    if len(text.strip()) < 20:
        raise HTTPException(
            status_code=422,
            detail="Could not extract meaningful text from this file. It may be a scanned/image-only PDF.",
        )

    return text
